"""Gera supabase/seed.sql a partir da planilha Amigos_do_Alface_26.xlsx.

Uso:  python3 scripts/gerar_seed.py caminho/da/planilha.xlsx [ano]

Lê cada aba de rodada (nome "DD_MM"): coluna D/F = Time BRANCO (nome/gols),
J/H = Time VERDE (nome/gols), linhas 30-33 = goleiros. Gols digitados sem nome
ao lado (gol contra / sem autor) contam no placar do time, como na planilha.
Também cadastra os atletas da aba "Atletas" que ainda não jogaram.
"""
import json
import re
import sys
import warnings

import openpyxl

warnings.filterwarnings('ignore')


def q(s):
    return "'" + str(s).replace("'", "''") + "'"


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else 'Amigos_do_Alface_26.xlsx'
    ano = int(sys.argv[2]) if len(sys.argv) > 2 else 2026
    wb = openpyxl.load_workbook(path)
    matches = []
    for ws in wb.worksheets:
        if not re.match(r'^\d\d_\d\d$', ws.title):
            continue
        d, m = ws.title.split('_')
        jogadores = []
        extra = {'branco': 0, 'verde': 0}
        for r in range(5, 34):
            for nc, gc, time in (('D', 'F', 'branco'), ('J', 'H', 'verde')):
                nome = ws[f'{nc}{r}'].value
                gols = int(ws[f'{gc}{r}'].value or 0)
                if isinstance(nome, str) and nome.strip() and nome.strip().upper() != 'GOLEIROS':
                    jogadores.append({'nome': ' '.join(nome.split()).upper(), 'time': time,
                                      'goleiro': r >= 30, 'gols': gols, 'amarelos': 0, 'vermelhos': 0})
                elif gols and r <= 27:
                    extra[time] += gols
        if jogadores:
            matches.append((f'{ano}-{m}-{d}', extra, jogadores))
    matches.sort()

    atletas = []
    if 'Atletas' in wb.sheetnames:
        ws = wb['Atletas']
        for r in range(5, ws.max_row + 1):
            n = ws[f'C{r}'].value
            if isinstance(n, str) and n.strip():
                atletas.append(' '.join(n.split()).upper())

    goleiros = sorted({j['nome'] for _, _, js in matches for j in js if j['goleiro']} - set(atletas))

    out = [
        '-- Carga inicial: rodadas da planilha Amigos do Alface (gerado por scripts/gerar_seed.py).',
        '-- Rode DEPOIS do schema.sql, no SQL Editor do Supabase. Pode rodar de novo:',
        '-- rodadas que já existem (mesma data) são ignoradas.',
        '',
    ]
    for nome in atletas:
        out.append(f'insert into public.players (nome) values ({q(nome)}) on conflict do nothing;')
    for nome in goleiros:
        out.append(f'insert into public.players (nome, goleiro) values ({q(nome)}, true) on conflict do nothing;')
    out.append('')
    for data, extra, js in matches:
        out.append(
            f"select public._save_match(null, {q(data)}::date, {extra['branco']}, {extra['verde']}, null, null, "
            f"{q(json.dumps(js, ensure_ascii=False))}::jsonb) "
            f"where not exists (select 1 from public.matches where data = {q(data)}::date);"
        )
    open('supabase/seed.sql', 'w', encoding='utf-8').write('\n'.join(out) + '\n')
    print(f'{len(matches)} rodadas, {len(atletas)} atletas, {len(goleiros)} goleiros -> supabase/seed.sql')


if __name__ == '__main__':
    main()
